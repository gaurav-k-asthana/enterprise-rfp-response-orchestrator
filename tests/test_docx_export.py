from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

from rfp_orchestrator.docx_export import (
    DOCUMENT_AUTHOR,
    DOCUMENT_SUBJECT,
    DOCUMENT_TITLE,
    FOOTER_TEXT,
    HEADER_TEXT,
    DocxExportError,
    generate_basic_response_docx,
)
from rfp_orchestrator.safety import SAFETY_NOTICE_BODY, SAFETY_NOTICE_TITLE
from rfp_orchestrator.sample_requirements import load_sample_requirements
from rfp_orchestrator.ui import stream_sample_requirement


def _paragraph_text(document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


def _evidence(chunk_id: str, *, text: str = "The capability is documented.") -> dict:
    return {
        "chunk_id": chunk_id,
        "doc_id": "DOC-001",
        "domain": "product",
        "title": "Product Capability Guide",
        "text": text,
        "version": "1.0",
        "effective_date": "2026-07-15",
        "authority_rank": 1,
        "source_status": "current",
        "score": 0.9,
        "retrieval_method": "dense",
    }


def _specialist_output(
    *,
    support_status: str = "SUPPORTED",
    claims: list[dict] | None = None,
) -> dict:
    return {
        "specialist": "product",
        "support_status": support_status,
        "proposed_answer": "The documented capability is supported.",
        "claims": claims
        or [
            {
                "claim_id": "product-claim-001",
                "text": "The capability is documented.",
                "supported": True,
                "evidence_ids": ["EVIDENCE-001"],
            }
        ],
    }


def test_finalized_state_generates_a_valid_requirement_response_docx() -> None:
    run = stream_sample_requirement(
        load_sample_requirements()[0],
        1,
        "step-3-15-finalized",
    )

    payload = generate_basic_response_docx(run.state)
    document = Document(BytesIO(payload))
    text = _paragraph_text(document)

    assert payload.startswith(b"PK")
    with ZipFile(BytesIO(payload)) as package:
        assert "word/document.xml" in package.namelist()
    assert DOCUMENT_TITLE in text
    assert "Requirement RFP-001 | Finalized" in text
    assert "Requirement" in text
    assert run.state["original_text"] in text
    assert "Final response" in text
    assert run.state["final_answer"] in text
    assert f"{SAFETY_NOTICE_TITLE}. {SAFETY_NOTICE_BODY}" in text
    assert "Evidence and support" in text
    assert "Product specialist — Supported" in text
    assert (
        "Aggregate support status: Supported. All atomic claims in this specialist "
        "output are supported."
    ) in text
    assert "Atomic claim product-claim-001 — Supported" in "\n".join(text)
    assert "Cited evidence" in text
    assert "Product Availability Matrix [PROD-AVAIL-001::chunk-001]" in text
    assert "Approval notes" in text
    assert (
        "No human approval was required or recorded for this requirement." in text
    )


@pytest.mark.parametrize(
    ("status", "expected"),
    [
        (
            "NEEDS_HUMAN",
            "awaiting human review. No final response has been authorized.",
        ),
        (
            "REJECTED",
            "rejected during human review. No final response was approved.",
        ),
        ("PENDING", "has not yet been processed. No final response is available."),
        ("IN_PROGRESS", "still being processed. No final response is available."),
    ],
)
def test_nonfinal_state_exports_a_truthful_status_instead_of_an_answer(
    status: str,
    expected: str,
) -> None:
    payload = generate_basic_response_docx(
        {
            "requirement_id": "RFP-TEST",
            "original_text": "Describe the supported deployment model.",
            "final_status": status,
            "final_answer": None,
        }
    )
    text = "\n".join(_paragraph_text(Document(BytesIO(payload))))

    assert "Current status" in text
    assert expected in text
    assert "Final response\n" not in text


def test_document_uses_the_locked_business_brief_geometry_and_furniture() -> None:
    document = Document(
        BytesIO(
            generate_basic_response_docx(
                {
                    "requirement_id": "RFP-TEST",
                    "original_text": "Confirm the documented capability.",
                    "final_status": "FINALIZED",
                    "final_answer": "The documented capability is supported.",
                }
            )
        )
    )
    section = document.sections[0]

    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.top_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.bottom_margin == Inches(1)
    assert section.left_margin == Inches(1)
    assert section.header_distance == Twips(708)
    assert section.footer_distance == Twips(708)
    assert section.header.paragraphs[0].text == HEADER_TEXT
    assert section.footer.paragraphs[0].text == FOOTER_TEXT
    assert document.styles["Normal"].font.name == "Calibri"
    assert document.styles["Normal"].font.size == Pt(11)
    assert document.styles["Heading 1"].font.size == Pt(16)
    assert document.styles["Safety Notice"].font.size == Pt(10)
    assert document.styles["Title"]._element.pPr.find(qn("w:pBdr")) is None


def test_document_metadata_is_synthetic_and_contains_no_user_identity() -> None:
    document = Document(
        BytesIO(
            generate_basic_response_docx(
                {
                    "requirement_id": "RFP-TEST",
                    "original_text": "Confirm the documented capability.",
                    "final_status": "FINALIZED",
                    "final_answer": "The documented capability is supported.",
                }
            )
        )
    )

    assert document.core_properties.title == DOCUMENT_TITLE
    assert document.core_properties.subject == DOCUMENT_SUBJECT
    assert document.core_properties.author == DOCUMENT_AUTHOR
    assert "Gaurav" not in document.core_properties.author


def test_empty_evidence_and_support_sections_are_explicit_not_invented() -> None:
    payload = generate_basic_response_docx(
        {
            "requirement_id": "RFP-TEST",
            "original_text": "Confirm the documented capability.",
            "final_status": "FINALIZED",
            "final_answer": "The documented capability is supported.",
        }
    )
    text = "\n".join(_paragraph_text(Document(BytesIO(payload))))

    assert "Evidence and support" in text
    assert (
        "No specialist support assessment or atomic claims were recorded for this path."
        in text
    )
    assert "Cited evidence" in text
    assert "No evidence records were cited for this path." in text
    assert "Approval notes" in text


def test_partial_aggregate_is_distinct_from_each_atomic_claim_boolean() -> None:
    state = {
        "requirement_id": "RFP-TEST",
        "original_text": "Confirm both documented capabilities.",
        "final_status": "FINALIZED",
        "final_answer": "One capability is supported and one is not documented.",
        "merged_specialist_outputs": [
            _specialist_output(
                support_status="PARTIAL",
                claims=[
                    {
                        "claim_id": "product-claim-001",
                        "text": "Capability A is documented.",
                        "supported": True,
                        "evidence_ids": ["EVIDENCE-001"],
                    },
                    {
                        "claim_id": "product-claim-002",
                        "text": "Capability B is not documented.",
                        "supported": False,
                        "evidence_ids": [],
                    },
                ],
            )
        ],
        "evidence": [_evidence("EVIDENCE-001")],
    }

    text = "\n".join(
        _paragraph_text(Document(BytesIO(generate_basic_response_docx(state))))
    )

    assert "Product specialist — Partial" in text
    assert "Some, but not all, atomic claims" in text
    assert "Atomic claim product-claim-001 — Supported" in text
    assert "Atomic claim product-claim-002 — Unsupported" in text
    assert "Citations: EVIDENCE-001" in text
    assert "Citations: None" in text


def test_citation_section_includes_only_evidence_actually_cited_by_a_claim() -> None:
    state = {
        "requirement_id": "RFP-TEST",
        "original_text": "Confirm the documented capability.",
        "final_status": "FINALIZED",
        "final_answer": "The documented capability is supported.",
        "merged_specialist_outputs": [_specialist_output()],
        "evidence": [
            _evidence("EVIDENCE-001"),
            _evidence("UNUSED-002", text="This retrieved item was not cited."),
        ],
    }

    text = "\n".join(
        _paragraph_text(Document(BytesIO(generate_basic_response_docx(state))))
    )

    assert "Product Capability Guide [EVIDENCE-001]" in text
    assert "UNUSED-002" not in text
    assert "This retrieved item was not cited." not in text


def test_every_claim_citation_must_resolve_to_saved_evidence() -> None:
    state = {
        "requirement_id": "RFP-TEST",
        "original_text": "Confirm the documented capability.",
        "final_status": "FINALIZED",
        "final_answer": "The documented capability is supported.",
        "merged_specialist_outputs": [_specialist_output()],
        "evidence": [],
    }

    with pytest.raises(DocxExportError, match="must resolve"):
        generate_basic_response_docx(state)


def test_aggregate_support_must_match_atomic_claim_booleans() -> None:
    state = {
        "requirement_id": "RFP-TEST",
        "original_text": "Confirm the documented capability.",
        "final_status": "FINALIZED",
        "final_answer": "The capability is not supported.",
        "merged_specialist_outputs": [
            _specialist_output(
                support_status="SUPPORTED",
                claims=[
                    {
                        "claim_id": "product-claim-001",
                        "text": "The capability is not documented.",
                        "supported": False,
                        "evidence_ids": [],
                    }
                ],
            )
        ],
    }

    with pytest.raises(DocxExportError, match="aggregate status"):
        generate_basic_response_docx(state)


def test_human_decision_history_is_rendered_in_auditable_order() -> None:
    state = {
        "requirement_id": "RFP-005",
        "original_text": "Commit to a customer-specific SLA.",
        "final_status": "FINALIZED",
        "final_answer": "The approved standard is offered without a new commitment.",
        "human_decision_history": [
            {
                "requirement_id": "RFP-005",
                "decision": "ADD_GUIDANCE",
                "reviewer": "reviewer@example.test",
                "timestamp": "2026-08-30T14:00:00Z",
                "review_reason": "ORGANIZATIONAL_AUTHORITY",
                "guidance": "Use only the approved standard.",
            },
            {
                "requirement_id": "RFP-005",
                "decision": "EDIT_AND_APPROVE",
                "reviewer": "reviewer@example.test",
                "timestamp": "2026-08-30T14:05:00Z",
                "review_reason": "ORGANIZATIONAL_AUTHORITY",
                "edited_answer": (
                    "The approved standard is offered without a new commitment."
                ),
            },
        ],
    }

    text = "\n".join(
        _paragraph_text(Document(BytesIO(generate_basic_response_docx(state))))
    )

    first = text.index("Decision 1 — Add Guidance")
    second = text.index("Decision 2 — Edit And Approve")
    assert first < second
    assert "Review reason: Organizational Authority" in text
    assert "Guidance: Use only the approved standard." in text
    assert "Edited answer: The approved standard is offered" in text


def test_awaiting_review_without_a_decision_is_labeled_truthfully() -> None:
    state = {
        "requirement_id": "RFP-005",
        "original_text": "Commit to a customer-specific SLA.",
        "final_status": "NEEDS_HUMAN",
        "final_answer": None,
        "awaiting_human_review": True,
    }

    text = "\n".join(
        _paragraph_text(Document(BytesIO(generate_basic_response_docx(state))))
    )

    assert "Awaiting human review; no decision has been recorded yet." in text


def test_approval_record_for_another_requirement_fails_closed() -> None:
    state = {
        "requirement_id": "RFP-005",
        "original_text": "Commit to a customer-specific SLA.",
        "final_status": "REJECTED",
        "final_answer": None,
        "approval": {
            "requirement_id": "RFP-OTHER",
            "decision": "REJECT",
            "reviewer": "reviewer@example.test",
            "timestamp": "2026-08-30T14:00:00Z",
        },
    }

    with pytest.raises(DocxExportError, match="another requirement"):
        generate_basic_response_docx(state)


@pytest.mark.parametrize(
    "state",
    [
        {},
        {
            "requirement_id": " ",
            "original_text": "Requirement",
            "final_status": "FINALIZED",
            "final_answer": "Answer",
        },
        {
            "requirement_id": "RFP-TEST",
            "original_text": "Requirement",
            "final_status": "UNKNOWN",
            "final_answer": None,
        },
        {
            "requirement_id": "RFP-TEST",
            "original_text": "Requirement",
            "final_status": "FINALIZED",
            "final_answer": None,
        },
        {
            "requirement_id": "RFP-TEST",
            "original_text": "Requirement",
            "final_status": "NEEDS_HUMAN",
            "final_answer": "Unsafe premature answer",
        },
    ],
)
def test_malformed_or_contradictory_export_state_fails_closed(state: dict) -> None:
    with pytest.raises(DocxExportError):
        generate_basic_response_docx(state)
