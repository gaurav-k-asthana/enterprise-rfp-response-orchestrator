"""Simple, reviewable requirement-level DOCX export for Phase 3."""

from __future__ import annotations

from collections.abc import Mapping
from io import BytesIO
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from pydantic import ValidationError

from rfp_orchestrator.models import (
    ApprovalDecision,
    RequirementStatus,
    SpecialistOutput,
)
from rfp_orchestrator.retrieval import EvidenceChunk
from rfp_orchestrator.safety import SAFETY_NOTICE_BODY, SAFETY_NOTICE_TITLE

DOCUMENT_TITLE = "Enterprise RFP Response"
DOCUMENT_SUBJECT = "Synthetic requirement-level RFP response"
DOCUMENT_AUTHOR = "Northstar Systems (synthetic demonstration)"
HEADER_TEXT = "Northstar Systems | Enterprise RFP Response"
FOOTER_TEXT = "Synthetic demonstration | Draft for human review"

INK = RGBColor(31, 41, 55)
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(91, 103, 117)
GOLD = RGBColor(122, 90, 0)
NOTICE_FILL = "FFF8E8"
NOTICE_BORDER = "C7A94A"
SUPPORTED_FILL = "EAF5ED"
PARTIAL_FILL = "FFF5D9"
UNSUPPORTED_FILL = "FDECEC"
DETAIL_BORDER = "D7DEE8"

SUPPORT_EXPLANATIONS = {
    "SUPPORTED": "All atomic claims in this specialist output are supported.",
    "PARTIAL": "Some, but not all, atomic claims in this specialist output are supported.",
    "UNSUPPORTED": "None of the atomic claims in this specialist output are supported.",
}


class DocxExportError(ValueError):
    """Raised when saved graph state cannot produce a truthful basic export."""


def _required_text(state: Mapping[str, Any], field: str) -> str:
    value = state.get(field)
    if not isinstance(value, str) or not value.strip():
        raise DocxExportError(f"DOCX export requires a nonblank {field}")
    return value.strip()


def _optional_text(value: object, field: str) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str) or not value.strip():
        raise DocxExportError(f"DOCX export requires nonblank {field} when present")
    return value.strip()


def _validated_content(
    state: Mapping[str, Any],
) -> tuple[str, str, RequirementStatus, str]:
    requirement_id = _required_text(state, "requirement_id")
    requirement_text = _required_text(state, "original_text")
    try:
        final_status = RequirementStatus(state.get("final_status"))
    except (TypeError, ValueError) as error:
        raise DocxExportError("DOCX export requires a recognized final status") from error

    raw_answer = state.get("final_answer")
    if final_status is RequirementStatus.FINALIZED:
        if not isinstance(raw_answer, str) or not raw_answer.strip():
            raise DocxExportError("a finalized DOCX export requires a final answer")
        response = raw_answer.strip()
    else:
        if raw_answer is not None:
            raise DocxExportError("a nonfinal DOCX export cannot contain a final answer")
        response = {
            RequirementStatus.NEEDS_HUMAN: (
                "This requirement is awaiting human review. No final response has "
                "been authorized."
            ),
            RequirementStatus.REJECTED: (
                "This requirement was rejected during human review. No final "
                "response was approved."
            ),
            RequirementStatus.PENDING: (
                "This requirement has not yet been processed. No final response is "
                "available."
            ),
            RequirementStatus.IN_PROGRESS: (
                "This requirement is still being processed. No final response is "
                "available."
            ),
        }[final_status]
    return requirement_id, requirement_text, final_status, response


def _set_font(run: Any, *, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def _configure_styles(document: DocumentObject) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = NAVY
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    title_properties = title._element.get_or_add_pPr()
    inherited_border = title_properties.find(qn("w:pBdr"))
    if inherited_border is not None:
        title_properties.remove(inherited_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.italic = False
    subtitle.font.color.rgb = MUTED
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.keep_with_next = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Safety Notice" not in styles:
        notice = styles.add_style("Safety Notice", WD_STYLE_TYPE.PARAGRAPH)
    else:
        notice = styles["Safety Notice"]
    notice.base_style = normal
    notice.font.name = "Calibri"
    notice.font.size = Pt(10)
    notice.font.color.rgb = INK
    notice._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    notice._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    notice.paragraph_format.space_before = Pt(4)
    notice.paragraph_format.space_after = Pt(14)
    notice.paragraph_format.left_indent = Inches(0.14)
    notice.paragraph_format.right_indent = Inches(0.14)
    notice.paragraph_format.line_spacing = 1.10

    detail_styles = {
        "Detail Label": (10, NAVY, True, 8, 2),
        "Detail Metadata": (9, MUTED, False, 0, 4),
        "Evidence Excerpt": (10, INK, False, 0, 10),
    }
    for name, (size, color, bold, before, after) in detail_styles.items():
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.base_style = normal
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_together = True


def _shade_and_border_notice(paragraph: Any) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), NOTICE_FILL)
    properties.append(shading)

    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), NOTICE_BORDER)
    borders.append(left)
    properties.append(borders)


def _shade_and_border_detail(paragraph: Any, *, fill: str = "F7F9FC") -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)

    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), DETAIL_BORDER)
    borders.append(bottom)
    properties.append(borders)


def _label(value: str) -> str:
    return value.replace("_", " ").title()


def _validated_specialist_outputs(
    state: Mapping[str, Any],
) -> list[SpecialistOutput]:
    raw_outputs = state.get("merged_specialist_outputs", [])
    if not isinstance(raw_outputs, list):
        raise DocxExportError("merged specialist outputs must be a list")

    outputs: list[SpecialistOutput] = []
    seen_specialists: set[str] = set()
    seen_claim_ids: set[str] = set()
    try:
        for raw_output in raw_outputs:
            output = SpecialistOutput.model_validate(raw_output)
            specialist = output.specialist.value
            if specialist in seen_specialists:
                raise DocxExportError("specialist outputs cannot repeat a specialist")
            seen_specialists.add(specialist)
            for claim in output.claims:
                if not claim.claim_id.strip() or not claim.text.strip():
                    raise DocxExportError("claim IDs and claim text cannot be blank")
                if claim.claim_id in seen_claim_ids:
                    raise DocxExportError("claim IDs must be unique across the response")
                seen_claim_ids.add(claim.claim_id)
                if len(claim.evidence_ids) != len(set(claim.evidence_ids)):
                    raise DocxExportError("a claim cannot repeat a citation ID")
                if any(not item.strip() for item in claim.evidence_ids):
                    raise DocxExportError("citation IDs cannot be blank")
            outputs.append(output)
    except ValidationError as error:
        raise DocxExportError(
            "specialist support state is malformed or its aggregate status does not "
            "match its atomic claims"
        ) from error
    return outputs


def _validated_cited_evidence(
    state: Mapping[str, Any],
    outputs: list[SpecialistOutput],
) -> list[EvidenceChunk]:
    cited_ids = {
        evidence_id
        for output in outputs
        for claim in output.claims
        for evidence_id in claim.evidence_ids
    }
    raw_evidence = state.get("evidence", [])
    if not isinstance(raw_evidence, list):
        raise DocxExportError("evidence state must be a list")

    evidence_by_id: dict[str, EvidenceChunk] = {}
    try:
        for raw_item in raw_evidence:
            if not isinstance(raw_item, Mapping):
                raise DocxExportError("evidence records must be mappings")
            raw_id = raw_item.get("chunk_id")
            if raw_id not in cited_ids:
                continue
            item = EvidenceChunk.model_validate(raw_item)
            existing = evidence_by_id.get(item.chunk_id)
            if existing is not None and existing != item:
                raise DocxExportError(
                    "duplicate citation IDs cannot describe different evidence"
                )
            evidence_by_id[item.chunk_id] = item
    except ValidationError as error:
        raise DocxExportError("cited evidence state is malformed") from error

    missing_ids = sorted(cited_ids - evidence_by_id.keys())
    if missing_ids:
        raise DocxExportError(
            "every claim citation must resolve to saved evidence: "
            + ", ".join(missing_ids)
        )
    return [evidence_by_id[item] for item in sorted(evidence_by_id)]


def _validated_approval_records(state: Mapping[str, Any]) -> list[dict[str, str]]:
    raw_history = state.get("human_decision_history", [])
    if not isinstance(raw_history, list):
        raise DocxExportError("human decision history must be a list")
    raw_records: list[object] = list(raw_history)
    if not raw_records and state.get("approval") is not None:
        raw_records = [state["approval"]]

    records: list[dict[str, str]] = []
    requirement_id = str(state.get("requirement_id", "")).strip()
    for raw_record in raw_records:
        if not isinstance(raw_record, Mapping):
            raise DocxExportError("approval records must be mappings")
        record_requirement_id = _optional_text(
            raw_record.get("requirement_id"), "approval requirement_id"
        )
        if record_requirement_id != requirement_id:
            raise DocxExportError("approval record belongs to another requirement")
        try:
            decision = ApprovalDecision(raw_record.get("decision"))
        except (TypeError, ValueError) as error:
            raise DocxExportError("approval record has an unknown decision") from error
        reviewer = _optional_text(raw_record.get("reviewer"), "approval reviewer")
        timestamp = _optional_text(raw_record.get("timestamp"), "approval timestamp")
        if reviewer is None or timestamp is None:
            raise DocxExportError("approval record requires reviewer and timestamp")
        edited_answer = _optional_text(
            raw_record.get("edited_answer"), "approval edited_answer"
        )
        guidance = _optional_text(raw_record.get("guidance"), "approval guidance")
        review_reason = _optional_text(
            raw_record.get("review_reason"), "approval review_reason"
        )
        if decision is ApprovalDecision.EDIT_AND_APPROVE and edited_answer is None:
            raise DocxExportError("EDIT_AND_APPROVE requires an edited answer")
        if decision is ApprovalDecision.ADD_GUIDANCE and guidance is None:
            raise DocxExportError("ADD_GUIDANCE requires guidance")
        records.append(
            {
                "decision": decision.value,
                "reviewer": reviewer,
                "timestamp": timestamp,
                "review_reason": review_reason or "Not recorded",
                "edited_answer": edited_answer or "None",
                "guidance": guidance or "None",
            }
        )
    return records


def _add_labeled_paragraph(
    document: DocumentObject,
    label: str,
    value: str,
    *,
    style: str = "Normal",
) -> Any:
    paragraph = document.add_paragraph(style=style)
    label_run = paragraph.add_run(f"{label}: ")
    _set_font(label_run)
    label_run.bold = True
    value_run = paragraph.add_run(value)
    _set_font(value_run)
    return paragraph


def _add_support_sections(
    document: DocumentObject,
    outputs: list[SpecialistOutput],
) -> None:
    document.add_heading("Evidence and support", level=1)
    if not outputs:
        document.add_paragraph(
            "No specialist support assessment or atomic claims were recorded for this path."
        )
        return

    for output in outputs:
        status = output.support_status.value
        document.add_heading(
            f"{_label(output.specialist.value)} specialist — {_label(status)}",
            level=2,
        )
        status_paragraph = _add_labeled_paragraph(
            document,
            "Aggregate support status",
            f"{_label(status)}. {SUPPORT_EXPLANATIONS[status]}",
            style="Detail Metadata",
        )
        status_fill = {
            "SUPPORTED": SUPPORTED_FILL,
            "PARTIAL": PARTIAL_FILL,
            "UNSUPPORTED": UNSUPPORTED_FILL,
        }[status]
        _shade_and_border_detail(status_paragraph, fill=status_fill)
        for claim in output.claims:
            claim_label = "Supported" if claim.supported else "Unsupported"
            paragraph = _add_labeled_paragraph(
                document,
                f"Atomic claim {claim.claim_id} — {claim_label}",
                claim.text,
                style="Detail Label",
            )
            paragraph.paragraph_format.keep_with_next = True
            citations = ", ".join(claim.evidence_ids) if claim.evidence_ids else "None"
            _add_labeled_paragraph(
                document,
                "Citations",
                citations,
                style="Detail Metadata",
            )


def _add_citation_section(
    document: DocumentObject,
    evidence: list[EvidenceChunk],
) -> None:
    document.add_heading("Cited evidence", level=1)
    if not evidence:
        document.add_paragraph("No evidence records were cited for this path.")
        return
    for item in evidence:
        heading = document.add_paragraph(style="Detail Label")
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(f"{item.title} [{item.chunk_id}]")
        _set_font(run)
        run.bold = True
        metadata = (
            f"{_label(item.domain.value)} | Version {item.version} | Effective "
            f"{item.effective_date} | {_label(item.source_status)} | "
            f"{_label(item.retrieval_method.value)} retrieval"
        )
        document.add_paragraph(metadata, style="Detail Metadata")
        excerpt = " ".join(item.text.split())
        if len(excerpt) > 320:
            excerpt = f"{excerpt[:319].rstrip()}…"
        paragraph = _add_labeled_paragraph(
            document,
            "Evidence excerpt",
            excerpt,
            style="Evidence Excerpt",
        )
        _shade_and_border_detail(paragraph)


def _add_approval_section(
    document: DocumentObject,
    state: Mapping[str, Any],
    records: list[dict[str, str]],
) -> None:
    if records:
        document.add_page_break()
    document.add_heading("Approval notes", level=1)
    if not records:
        if state.get("awaiting_human_review") is True:
            document.add_paragraph(
                "Awaiting human review; no decision has been recorded yet."
            )
        else:
            document.add_paragraph(
                "No human approval was required or recorded for this requirement."
            )
        return

    for index, record in enumerate(records, start=1):
        document.add_heading(
            f"Decision {index} — {_label(record['decision'])}",
            level=2,
        )
        _add_labeled_paragraph(document, "Reviewer", record["reviewer"])
        _add_labeled_paragraph(document, "Timestamp", record["timestamp"])
        _add_labeled_paragraph(
            document,
            "Review reason",
            _label(record["review_reason"])
            if record["review_reason"] != "Not recorded"
            else record["review_reason"],
        )
        _add_labeled_paragraph(document, "Edited answer", record["edited_answer"])
        _add_labeled_paragraph(document, "Guidance", record["guidance"])


def _configure_page(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Twips(708)
    section.footer_distance = Twips(708)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run(HEADER_TEXT)
    _set_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run(FOOTER_TEXT)
    _set_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def build_basic_response_document(state: Mapping[str, Any]) -> DocumentObject:
    """Build the Step 3.16 response with evidence and approval provenance."""

    if not isinstance(state, Mapping):
        raise DocxExportError("DOCX export requires saved requirement state")
    requirement_id, requirement_text, final_status, response = _validated_content(
        state
    )
    outputs = _validated_specialist_outputs(state)
    evidence = _validated_cited_evidence(state, outputs)
    approval_records = _validated_approval_records(state)

    document = Document()
    _configure_styles(document)
    _configure_page(document)
    properties = document.core_properties
    properties.title = DOCUMENT_TITLE
    properties.subject = DOCUMENT_SUBJECT
    properties.author = DOCUMENT_AUTHOR
    properties.keywords = "synthetic RFP response, human review"

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("SYNTHETIC ENTERPRISE DEMONSTRATION")
    _set_font(run)
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = GOLD

    document.add_paragraph(DOCUMENT_TITLE, style="Title")
    status_label = final_status.value.replace("_", " ").title()
    document.add_paragraph(
        f"Requirement {requirement_id} | {status_label}",
        style="Subtitle",
    )

    notice = document.add_paragraph(style="Safety Notice")
    title_run = notice.add_run(f"{SAFETY_NOTICE_TITLE}. ")
    _set_font(title_run)
    title_run.bold = True
    body_run = notice.add_run(SAFETY_NOTICE_BODY)
    _set_font(body_run)
    _shade_and_border_notice(notice)

    document.add_heading("Requirement", level=1)
    document.add_paragraph(requirement_text)

    response_heading = (
        "Final response"
        if final_status is RequirementStatus.FINALIZED
        else "Current status"
    )
    document.add_heading(response_heading, level=1)
    document.add_paragraph(response)
    _add_support_sections(document, outputs)
    _add_citation_section(document, evidence)
    _add_approval_section(document, state, approval_records)
    return document


def generate_basic_response_docx(state: Mapping[str, Any]) -> bytes:
    """Serialize one guarded requirement response to DOCX bytes."""

    output = BytesIO()
    build_basic_response_document(state).save(output)
    return output.getvalue()
